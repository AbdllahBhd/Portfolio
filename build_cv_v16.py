"""Build Abdellah Bouhaddi's detailed UAE-focused CV v16."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Abdellah_Bouhaddi_Full_Stack_Software_Engineer_CV_v16.docx"

BLACK = RGBColor(0x00, 0x00, 0x00)
BODY = RGBColor(0x20, 0x27, 0x2B)
MUTED = RGBColor(0x55, 0x5F, 0x66)
ACCENT = RGBColor(0x0B, 0x67, 0x63)
FONT = "Aptos"


def set_font(run, size=10.2, bold=False, italic=False, color=BODY):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def format_paragraph(paragraph, before=0, after=2.7, line=1.06, keep_next=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.keep_with_next = keep_next
    fmt.widow_control = True


def add_hyperlink(paragraph, text, url, size=9.5):
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{key}"), FONT)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0B6763")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size_el = OxmlElement("w:sz")
    size_el.set(qn("w:val"), str(round(size * 2)))
    props.extend([fonts, color, underline, size_el])
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.extend([props, text_el])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_section(document, text):
    p = document.add_paragraph(style="Heading 1")
    format_paragraph(p, before=8, after=3.5, line=1.0, keep_next=True)
    set_font(p.add_run(text.upper()), 11.1, bold=True, color=BLACK)
    return p


def add_text(document, text, size=10.2, after=3.0):
    p = document.add_paragraph()
    format_paragraph(p, after=after, line=1.08)
    set_font(p.add_run(text), size=size)
    return p


def add_skill(document, label, text):
    p = document.add_paragraph()
    format_paragraph(p, after=1.7, line=1.03)
    set_font(p.add_run(label + ": "), 9.8, bold=True, color=BLACK)
    set_font(p.add_run(text), 9.8)
    return p


def add_role(document, title, dates, company, location):
    p = document.add_paragraph(style="Heading 2")
    format_paragraph(p, before=6.5, after=0.6, line=1.0, keep_next=True)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.05), 2)
    set_font(p.add_run(title), 10.6, bold=True, color=BLACK)
    p.add_run("\t")
    set_font(p.add_run(dates), 9.2, bold=True, color=ACCENT)

    meta = document.add_paragraph()
    format_paragraph(meta, after=1.8, line=1.0, keep_next=True)
    set_font(meta.add_run(company), 9.7, bold=True, color=ACCENT)
    set_font(meta.add_run("  |  " + location), 9.4, italic=True, color=MUTED)
    return p


def add_subheading(document, text):
    p = document.add_paragraph(style="Heading 2")
    format_paragraph(p, before=4.5, after=1.5, line=1.0, keep_next=True)
    set_font(p.add_run(text), 9.9, bold=True, color=BLACK)
    return p


def add_bullet(document, text):
    p = document.add_paragraph(style="List Bullet")
    format_paragraph(p, after=2.0, line=1.06)
    set_font(p.add_run(text), 10.0)
    return p


def page_break(document):
    p = document.add_paragraph()
    format_paragraph(p, after=0, line=1.0)
    p.add_run().add_break(WD_BREAK.PAGE)


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    format_paragraph(p, after=0, line=1.0)
    set_font(p.add_run("ABDELLAH BOUHADDI  |  "), 8.1, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "555F66")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "16")
    props.extend([color, size])
    run.extend([props, begin, instruction, separate, value, end])
    p._p.append(run)


def configure_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = BODY
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)

    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[name]
        style.font.name = FONT
        style.font.color.rgb = BLACK
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)

    bullets = document.styles["List Bullet"]
    bullets.font.name = FONT
    bullets.font.size = Pt(10)
    bullets._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    bullets._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    bullets.paragraph_format.left_indent = Inches(0.28)
    bullets.paragraph_format.first_line_indent = Inches(-0.16)
    bullets.paragraph_format.tab_stops.add_tab_stop(Inches(0.28))


def build():
    document = Document()
    configure_styles(document)

    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.52)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)
    add_page_number(section)

    name = document.add_paragraph(style="Title")
    format_paragraph(name, after=0.6, line=1.0)
    set_font(name.add_run("ABDELLAH BOUHADDI"), 21, bold=True, color=BLACK)

    title = document.add_paragraph()
    format_paragraph(title, after=2.7, line=1.0)
    set_font(title.add_run("FULL STACK SOFTWARE ENGINEER"), 11.4, bold=True, color=BLACK)
    set_font(title.add_run("  |  C#  ASP.NET CORE  REACT"), 10.2, bold=True, color=ACCENT)

    contact = document.add_paragraph()
    format_paragraph(contact, after=0.7, line=1.0)
    set_font(contact.add_run("Dubai, UAE  |  +971 58 534 2499  |  douibouhaddi@gmail.com"), 9.7)

    links = document.add_paragraph()
    format_paragraph(links, after=5, line=1.0)
    add_hyperlink(links, "Portfolio", "https://abdllahbhd.github.io/Portfolio")
    set_font(links.add_run("  |  "), 9.3, color=MUTED)
    add_hyperlink(links, "LinkedIn", "https://www.linkedin.com/in/abdellah-bouhaddi-1688421a9/")
    set_font(links.add_run("  |  "), 9.3, color=MUTED)
    add_hyperlink(links, "GitHub", "https://github.com/AbdllahBhd")

    add_section(document, "Professional Profile")
    add_text(
        document,
        "Full stack software engineer with 3+ years of experience building and improving business applications in C# and ASP.NET Core, React and Razor Pages, SQL, payments, and production delivery. At MATMAK in Dubai, contribute to an established .NET SaaS with the engineering team and deliver client applications from implementation through launch. Combine application development with Linux infrastructure, AWS, CI/CD, monitoring, and direct collaboration with management and clients. Earlier experience includes Java Spring Boot microservices, Flutter applications, and computer science instruction.",
        size=10.2,
        after=3.2,
    )

    add_section(document, "Technical Skills")
    add_skill(document, "Backend", "C#, ASP.NET Core 6 and 8, ASP.NET Web API, RESTful APIs, Entity Framework Core, LINQ, dependency injection, JWT and OAuth, Swagger and OpenAPI, OOP, SOLID principles, design patterns")
    add_skill(document, "Frontend and mobile", "React, Razor Pages, JavaScript, TypeScript, HTML5, CSS3, Bootstrap, responsive UI, Flutter, Dart, Android, Firebase")
    add_skill(document, "Data", "SQL Server, PostgreSQL, MySQL, MongoDB, Firestore, schema design, query optimization, relational and NoSQL data")
    add_skill(document, "Cloud and delivery", "AWS EC2, Linux, Docker, Nginx, SSL and TLS, GitHub Actions, Jenkins, CI/CD, staging, backups, DNS, Git")
    add_skill(document, "Integrations and quality", "Stripe, Network International, webhooks, xUnit, JUnit, unit and integration testing, Prometheus, Grafana, Loki, Promtail")
    add_skill(document, "Additional technologies", "Java, Spring Boot, Hibernate, PHP, Laravel, Filament, WordPress, WooCommerce, Angular, Jira, Agile and Scrum")

    add_section(document, "Professional Experience")
    add_role(document, "Full Stack Developer", "June 2024 - Present", "MATMAK", "Dubai, UAE")
    add_subheading(document, "ASP.NET Core SaaS and product engineering")
    for item in (
        "Contribute with the engineering team to an established ASP.NET Core SaaS and PPF cutting platform, extending REST APIs, Razor Pages, and workflows for authentication, subscriptions, invoicing, pricing, credits, and cutting operations.",
        "Work within a modular .NET solution spanning API, Client, Admin, Models, and Services projects, using Entity Framework Core and SQL-backed application services.",
        "Redesigned customer-facing pages and the administration interface, connecting responsive frontend changes to application services and data workflows.",
        "Integrated Stripe for one-time payments, recurring subscriptions, invoicing, and webhook-based status updates.",
        "Improved affected EF Core and database read and write operations by about 40% through query tuning and schema redesign.",
        "Collaborate directly with management and teammates on requirements, priorities, troubleshooting, code review, production releases, and ongoing support.",
    ):
        add_bullet(document, item)

    page_break(document)
    add_section(document, "MATMAK Delivery and Infrastructure")
    add_subheading(document, "Cloud infrastructure and production operations")
    for item in (
        "Provision Linux cloud and VPS servers from purchase through go-live, including operating system setup, SSH access, security hardening, firewalls, Nginx reverse proxy, SSL and TLS, and application runtime deployment.",
        "Manage AWS EC2, staging environments, automated backups, domains, subdomains, DNS, and Google Workspace email for application and client delivery.",
        "Build deployment workflows with GitHub Actions and monitor production services with Prometheus, Grafana, Loki, and Promtail.",
    ):
        add_bullet(document, item)

    add_subheading(document, "XSHIELD ecommerce platform  Production")
    for item in (
        "Delivered the public WooCommerce store at xshield-protect.com from infrastructure setup through production launch, covering Linux hosting, domain and DNS, Google Workspace, deployment, and operational support.",
        "Built custom pages, PHP customizations, API integrations, and WordPress plugins, including import and export workflows beyond standard WooCommerce configuration.",
        "Configured Stripe and Network International payment gateways for UAE ecommerce checkout.",
    ):
        add_bullet(document, item)

    add_subheading(document, "HEXIS Training Centre  Staging")
    for item in (
        "Architected and built a multilingual training application with a Laravel backend, Filament administration, and React public interface.",
        "Implemented course planning, training locations, registrations, certifications, transactional email, and the Hex'pert quiz flow with emailed access, automatic submission, scoring, and results.",
        "Used Cursor, OpenAI Codex, custom agent skills, and Model Context Protocol tools where useful, while reviewing and validating changes before release.",
    ):
        add_bullet(document, item)

    add_subheading(document, "HEXIS PRECUT LAB and BODYFENCE  Production")
    for item in (
        "Developed and enhanced a custom WordPress and PHP plugin for bodyfence.net/hpl that integrates with a dedicated MATMAK API instance.",
        "Extended pages and forms with custom shortcodes that fetch and display vehicle pattern data and car kits from the API.",
        "Maintain additional WordPress sites with attention to performance, security, usability, hosting, and production support.",
    ):
        add_bullet(document, item)

    page_break(document)
    add_section(document, "Earlier Experience")
    add_role(document, "Mobile Application Developer", "April 2024 - June 2024", "Freelance", "Dubai, UAE")
    for item in (
        "Developed cross-platform Android and iOS applications with Flutter and Dart, responsive Material interfaces, and custom navigation flows.",
        "Integrated Firebase Authentication, Cloud Firestore, Cloud Storage, and push notifications.",
    ):
        add_bullet(document, item)

    add_role(document, "Backend Developer  Java Spring Boot", "September 2023 - April 2024", "Tech-Instinct", "Bejaia, Algeria")
    for item in (
        "Maintained an existing Spring Boot production codebase and REST microservices for the RH Partner HR and payroll SaaS, with each service deployed in a separate Docker container.",
        "Designed database schemas and optimized SQL queries, improving affected query performance by 35% and supporting application scalability.",
        "Wrote JUnit unit and integration tests with more than 90% coverage in the tested scope, and diagnosed production defects to support reliable releases.",
    ):
        add_bullet(document, item)

    add_role(document, "Part Time Computer Science Instructor", "September 2022 - June 2023", "ESTIN and University Abderrahmane Mira", "Bejaia, Algeria")
    add_bullet(document, "Mentored more than 100 students and created practical labs in programming, algorithms and data structures, computer networks, web development, and mobile development.")

    add_section(document, "Education")
    add_role(document, "Master's Degree in Software Engineering", "2020 - 2022", "University Abderrahmane Mira", "Bejaia, Algeria")
    add_bullet(document, "Built a MERN progressive web ecommerce application with real-time inventory management, payment processing, and a responsive checkout flow.")
    add_role(document, "Bachelor's Degree in Computer Science", "2017 - 2020", "University Abderrahmane Mira", "Bejaia, Algeria")
    add_bullet(document, "Built Java inventory management software and desktop C++ and Android game implementations using the Minimax algorithm; participated in the Computer Science Club.")

    add_section(document, "Certifications and Languages")
    add_skill(document, "Certifications", "Java Developer Certificate, January to March 2020; Huawei Certification in Cloud Computing")
    add_skill(document, "Languages", "Kabyle native; Arabic, English, and French fluent")

    add_section(document, "Selected Work")
    for label, url in (
        ("Portfolio", "https://abdllahbhd.github.io/Portfolio"),
        ("XSHIELD", "https://xshield-protect.com/"),
        ("HEXIS PRECUT LAB and BODYFENCE", "https://www.bodyfence.net/en/hpl/"),
        ("RH Partner", "https://rh-partner.com/"),
    ):
        p = document.add_paragraph()
        format_paragraph(p, after=1.4, line=1.0)
        set_font(p.add_run(label + ": "), 9.5, bold=True, color=BLACK)
        add_hyperlink(p, url, url, size=9.4)

    document.core_properties.title = "Abdellah Bouhaddi Full Stack Software Engineer CV"
    document.core_properties.subject = "Full stack and .NET software engineering roles in Dubai and the UAE"
    document.core_properties.author = "Abdellah Bouhaddi"
    document.core_properties.keywords = (
        "C#, ASP.NET Core, .NET, Web API, REST, React, Razor Pages, Entity Framework Core, "
        "SQL Server, PostgreSQL, AWS, Linux, Docker, CI/CD, GitHub Actions, Dubai, UAE"
    )
    document.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()

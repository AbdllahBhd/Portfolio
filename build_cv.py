"""Build the Abdellah Bouhaddi CV (A4, .NET full stack roles in the UAE) as matching DOCX and PDF.

Every run creates a new version: the number in CV_VERSION is incremented, the numbered
DOCX and PDF are written, and the PDF is copied to the portfolio site assets so the next
push to main publishes it. Use --version N to build a specific number instead.
"""

from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate, Table, TableStyle


ROOT = Path(__file__).resolve().parent
VERSION_FILE = ROOT / "CV_VERSION"  # tracked in git so the counter is the same on every machine
BASENAME = "Abdellah_Bouhaddi_Full_Stack_Software_Engineer_CV"
# The portfolio site links to this fixed filename; GitHub Pages serves it after a push to main.
SITE_PDF = ROOT / "portfolio" / "assets" / "Abdellah_Bouhaddi_CV.pdf"


def next_version(forced: int | None) -> int:
    if forced is not None:
        return forced
    current = int(VERSION_FILE.read_text(encoding="utf-8").strip() or 0)
    return current + 1


def output_paths(version: int) -> tuple[Path, Path]:
    docx_out = ROOT / f"{BASENAME}_v{version}.docx"
    pdf_out = ROOT / "output" / "pdf" / f"{BASENAME}_v{version}.pdf"
    return docx_out, pdf_out



FONT = "Aptos"
BLACK = RGBColor(0x11, 0x16, 0x1A)
BODY = RGBColor(0x26, 0x2D, 0x33)
MUTED = RGBColor(0x55, 0x5E, 0x66)
ACCENT = RGBColor(0x0B, 0x4F, 0x6C)
ACCENT_HEX = "0B4F6C"
PDF_BLACK = colors.HexColor("#11161A")
PDF_BODY = colors.HexColor("#262D33")
PDF_MUTED = colors.HexColor("#555E66")
PDF_ACCENT = colors.HexColor("#" + ACCENT_HEX)

MARGIN = 0.6  # inches, left and right
CONTENT_WIDTH = 8.27 - 2 * MARGIN
PDF_WIDTH = CONTENT_WIDTH * 72 - 12  # points; reportlab frames pad 6pt per side

NAME = "ABDELLAH BOUHADDI"
HEADLINE = "FULL STACK SOFTWARE ENGINEER"
HEADLINE_STACK = ".NET  ·  C#  ·  ASP.NET CORE  ·  SQL  ·  AWS"
LOCATION = "Dubai, UAE"
PHONE = "+971 58 534 2499"
EMAIL = "douibouhaddi@gmail.com"
LINKS = [
    ("abdllahbhd.github.io/Portfolio", "https://abdllahbhd.github.io/Portfolio"),
    ("linkedin.com/in/abdellah-bouhaddi-1688421a9", "https://www.linkedin.com/in/abdellah-bouhaddi-1688421a9/"),
    ("github.com/AbdllahBhd", "https://github.com/AbdllahBhd"),
]

PROFILE = (
    "Full stack software engineer with 3+ years of experience delivering production business applications "
    "with C#, ASP.NET Core, Entity Framework Core, SQL, and Razor Pages. At MATMAK in Dubai, I extend an "
    "established ASP.NET Core SaaS platform, covering REST APIs, Stripe subscription billing, and customer and "
    "admin interfaces, and I own the full delivery chain: architecture, production databases, Linux and AWS "
    "servers, domains and DNS, CI/CD, monitoring, and production support. Improved affected database read and "
    "write performance by about 40% and lead client projects from requirements to launch. Earlier built Java "
    "Spring Boot microservices for an HR and payroll SaaS. Comfortable working directly with management and "
    "clients in English, Arabic, and French."
)

SKILLS = [
    ("Backend", "C#, ASP.NET Core, ASP.NET Web API, RESTful APIs, Entity Framework Core, LINQ, dependency injection, JWT and OAuth, Swagger and OpenAPI, OOP, SOLID, design patterns"),
    ("Frontend", "React, JavaScript, TypeScript, Razor Pages, HTML5, CSS3, Bootstrap, responsive UI"),
    ("Databases", "SQL Server, PostgreSQL, MySQL, MongoDB, Firestore; schema design, query optimization, stored procedures"),
    ("Cloud and DevOps", "AWS EC2, Linux, Docker, Nginx, SSL and TLS, GitHub Actions, Jenkins, CI/CD, Git, Prometheus, Grafana, Loki, Promtail"),
    ("Payments and testing", "Stripe (subscriptions, invoicing, webhooks), Network International, xUnit, JUnit, unit and integration testing"),
    ("Also experienced in", "Java, Spring Boot, Hibernate, microservices, PHP, Laravel, Filament, WordPress, WooCommerce, Flutter, Dart, Firebase, Angular"),
    ("Ways of working", "Agile and Scrum, Jira, code review, secure coding, AI-assisted development (Cursor, OpenAI Codex, MCP) with review and testing before release"),
]

MATMAK_INTRO = (
    "Dubai software company building an ASP.NET Core SaaS and PPF (paint protection film) cutting platform, "
    "plus web solutions for automotive clients."
)

MATMAK_PLATFORM = [
    "Extend an established ASP.NET Core SaaS and PPF cutting platform, building REST APIs, Razor Pages, and business "
    "workflows for authentication, subscriptions, invoicing, pricing, credits, and cutting operations.",
    "Work in a modular multi-project .NET solution (API, Client, Admin, Models, Services) using Entity Framework Core, "
    "LINQ, and SQL-backed application services.",
    "Redesigned the customer-facing pages and administration interface from a dated layout into a modern, responsive UI "
    "connected to existing application services and data workflows.",
    "Integrated Stripe one-time payments, recurring subscriptions, invoicing, and webhook-driven payment status updates.",
    "Improved affected read and write operations by about 40% by tuning EF Core queries and redesigning the database "
    "schema to remove redundant tables and costly joins.",
]

MATMAK_DELIVERY = [
    "Provision and harden Linux cloud and VPS servers from purchase to go-live: SSH access, firewalls, Nginx reverse "
    "proxy, SSL and TLS, application runtimes, and deployment of ASP.NET Core APIs and web applications.",
    "Manage AWS EC2, production databases, automated backups, and separate staging environments; automate builds and "
    "deployments with GitHub Actions.",
    "Built production monitoring with Prometheus, Grafana, Loki, and Promtail, with dashboards and alerts covering "
    "servers, APIs, Nginx, and application logs.",
    "Administer the company's full infrastructure estate: cloud accounts, servers, production databases, domains, DNS, "
    "email, and payment services, with responsibility for availability, security, backups, and access control.",
    "Lead new client projects from requirements and architecture through launch and support, working directly with "
    "management and clients on priorities, technical decisions, and troubleshooting.",
]

CLIENT_PROJECTS = [
    (
        "XSHIELD ecommerce",
        "WooCommerce, PHP  ·  Production",
        ("xshield-protect.com", "https://xshield-protect.com/"),
        [
            "Delivered the public store end to end: Linux hosting, domain and DNS, Google Workspace email, deployment, and operational support.",
            "Built custom PHP pages and plugins, API integrations, and import and export workflows; configured Stripe and Network International payment gateways for UAE checkout.",
        ],
    ),
    (
        "HEXIS Training Centre, France",
        "Laravel, Filament, React  ·  Staging",
        None,
        [
            "Architected and built a multilingual training platform with a Laravel backend, Filament administration, and a React public interface.",
            "Implemented course planning, locations, registrations, certifications, transactional email, and a quiz flow with emailed access, automatic submission, scoring, and results.",
        ],
    ),
    (
        "HEXIS PRECUT LAB / BODYFENCE",
        "WordPress, PHP, REST API  ·  Production",
        ("bodyfence.net/hpl", "https://www.bodyfence.net/en/hpl/"),
        [
            "Developed a custom WordPress plugin that consumes a dedicated MATMAK API through shortcodes to display vehicle patterns and car kits.",
        ],
    ),
]

EARLIER_ROLES = [
    (
        "Mobile Application Developer", "Apr 2024 – Jun 2024", "Freelance", "Dubai, UAE",
        [
            "Developed cross-platform Android and iOS applications with Flutter and Dart, integrating Firebase Authentication, Cloud Firestore, Cloud Storage, and push notifications.",
        ],
    ),
    (
        "Backend Developer, Java Spring Boot", "Sep 2023 – Apr 2024", "Tech-Instinct", "Bejaia, Algeria",
        [
            "Maintained and extended the production Spring Boot REST microservices of the RH Partner HR and payroll SaaS, each deployed in its own Docker container.",
            "Designed database schemas and optimized SQL queries, improving affected query performance by 35%.",
            "Wrote JUnit unit and integration tests reaching over 90% coverage in the tested scope, and resolved production defects to support reliable releases.",
        ],
    ),
    (
        "Part-Time Computer Science Instructor", "Sep 2022 – Jun 2023",
        "ESTIN Higher School of Computer Science and Technology, University Abderrahmane Mira", "Bejaia, Algeria",
        [
            "Taught and mentored 100+ students in programming, algorithms and data structures, networks, web, and mobile development; created practical lab material.",
        ],
    ),
]

EDUCATION = [
    (
        "Master's Degree in Software Engineering", "2020 – 2022", "University Abderrahmane Mira", "Bejaia, Algeria",
        "Thesis: MERN (MongoDB, Express, React, Node.js) progressive web ecommerce application with real-time inventory, payment processing, and a responsive checkout.",
    ),
    (
        "Bachelor's Degree in Computer Science", "2017 – 2020", "University Abderrahmane Mira", "Bejaia, Algeria",
        "Projects: Java inventory management software; C++ desktop and Android game using the Minimax algorithm.",
    ),
]

CERTIFICATIONS = "Java Developer Certificate (2020); Huawei Certification in Cloud Computing"
LANGUAGES = "English, Arabic, French (fluent); Kabyle (native)"


# ---------------------------------------------------------------- DOCX

def set_word_font(run, size=9.6, bold=False, italic=False, color=BODY):
    run.font.name = FONT
    props = run._element.get_or_add_rPr()
    fonts = props.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        props.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def word_format(paragraph, before=0, after=1.6, line=1.03, keep_next=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.keep_with_next = keep_next
    fmt.keep_together = True
    fmt.widow_control = True


def word_hyperlink(paragraph, text, url, size=8.8):
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{key}"), FONT)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT_HEX)
    size_el = OxmlElement("w:sz")
    size_el.set(qn("w:val"), str(round(size * 2)))
    props.extend([fonts, color, size_el])
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.extend([props, text_el])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def word_section(document, text):
    p = document.add_paragraph(style="Heading 1")
    word_format(p, before=7, after=2.5, line=1.0, keep_next=True)
    border = OxmlElement("w:bottom")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "6")
    border.set(qn("w:space"), "2")
    border.set(qn("w:color"), ACCENT_HEX)
    borders = OxmlElement("w:pBdr")
    borders.append(border)
    p._p.get_or_add_pPr().append(borders)
    set_word_font(p.add_run(text.upper()), 10.5, bold=True, color=ACCENT)


def word_role(document, title, dates, company, location):
    p = document.add_paragraph(style="Heading 2")
    word_format(p, before=4, after=0.3, line=1.0, keep_next=True)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(CONTENT_WIDTH), WD_TAB_ALIGNMENT.RIGHT)
    set_word_font(p.add_run(title), 10.2, bold=True, color=BLACK)
    set_word_font(p.add_run("\t" + dates), 9.0, bold=True, color=ACCENT)
    meta = document.add_paragraph()
    word_format(meta, after=1.2, line=1.0, keep_next=True)
    set_word_font(meta.add_run(company), 9.3, bold=True, color=BODY)
    set_word_font(meta.add_run("  |  " + location), 9.0, italic=True, color=MUTED)


def word_note(document, text):
    p = document.add_paragraph()
    word_format(p, after=1.6, line=1.0, keep_next=True)
    set_word_font(p.add_run(text), 9.0, italic=True, color=MUTED)


def word_subheading(document, title, detail=None, link=None, before=3.5):
    p = document.add_paragraph(style="Heading 3")
    word_format(p, before=before, after=0.8, line=1.0, keep_next=True)
    set_word_font(p.add_run(title), 9.7, bold=True, color=BLACK)
    if detail:
        set_word_font(p.add_run("  |  " + detail), 8.9, italic=True, color=MUTED)
    if link:
        set_word_font(p.add_run("  |  "), 8.9, color=MUTED)
        word_hyperlink(p, link[0], link[1], 8.9)


def word_bullet(document, text, keep_next=False):
    p = document.add_paragraph(style="List Bullet")
    word_format(p, after=1.3, line=1.03, keep_next=keep_next)
    set_word_font(p.add_run(text), 9.5)


def word_bullets(document, items):
    """Bullets of one short entry, kept on the same page."""
    for i, text in enumerate(items):
        word_bullet(document, text, keep_next=i < len(items) - 1)


def word_skill(document, label, text):
    p = document.add_paragraph()
    word_format(p, after=1.1, line=1.02)
    p.paragraph_format.left_indent = Inches(1.5)
    p.paragraph_format.first_line_indent = Inches(-1.5)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(1.5))
    set_word_font(p.add_run(label), 9.3, bold=True, color=BLACK)
    set_word_font(p.add_run("\t" + text), 9.3)


def word_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    word_format(p, after=0, line=1.0)
    set_word_font(p.add_run("Abdellah Bouhaddi  |  Page "), 7.5, color=MUTED)
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText"); instruction.set(qn("xml:space"), "preserve"); instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t"); value.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run = OxmlElement("w:r"); props = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "555E66")
    size = OxmlElement("w:sz"); size.set(qn("w:val"), "15")
    props.extend([color, size]); run.extend([props, begin, instruction, separate, value, end]); p._p.append(run)


def configure_word_styles(document):
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet"):
        style = document.styles[name]
        style.font.name = FONT
        style.font.color.rgb = BODY if name in ("Normal", "List Bullet") else BLACK
        style.font.italic = False
        rfonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn(f"w:{key}"), FONT)
    document.styles["Normal"].font.size = Pt(9.5)
    bullets = document.styles["List Bullet"]
    bullets.font.size = Pt(9.5)
    bullets.paragraph_format.left_indent = Inches(0.22)
    bullets.paragraph_format.first_line_indent = Inches(-0.14)


def build_docx():
    document = Document()
    configure_word_styles(document)
    section = document.sections[0]
    section.page_width = Inches(8.27); section.page_height = Inches(11.69)
    section.top_margin = Inches(0.5); section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(MARGIN); section.right_margin = Inches(MARGIN)
    section.header_distance = Inches(0.25); section.footer_distance = Inches(0.25)
    word_footer(section)

    name = document.add_paragraph(style="Title")
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    word_format(name, after=1, line=1.0)
    set_word_font(name.add_run(NAME), 21, bold=True, color=BLACK)
    # The default Title style carries a bottom border; override it on the paragraph.
    borders = OxmlElement("w:pBdr")
    none = OxmlElement("w:bottom")
    none.set(qn("w:val"), "nil")
    borders.append(none)
    name._p.get_or_add_pPr().append(borders)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    word_format(title, after=3, line=1.0)
    set_word_font(title.add_run(HEADLINE), 11, bold=True, color=ACCENT)
    set_word_font(title.add_run("   |   " + HEADLINE_STACK), 10, bold=True, color=BODY)

    contact = document.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    word_format(contact, after=1, line=1.0)
    set_word_font(contact.add_run(f"{LOCATION}  |  {PHONE}  |  "), 9.0)
    word_hyperlink(contact, EMAIL, "mailto:" + EMAIL, 9.0)

    links = document.add_paragraph()
    links.alignment = WD_ALIGN_PARAGRAPH.CENTER
    word_format(links, after=2, line=1.0)
    for i, (text, url) in enumerate(LINKS):
        if i:
            set_word_font(links.add_run("  |  "), 8.8, color=MUTED)
        word_hyperlink(links, text, url, 8.8)

    word_section(document, "Professional Summary")
    p = document.add_paragraph()
    word_format(p, after=1.5, line=1.04)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_word_font(p.add_run(PROFILE), 9.6)

    word_section(document, "Technical Skills")
    for label, text in SKILLS:
        word_skill(document, label, text)

    word_section(document, "Professional Experience")
    word_role(document, "Full Stack Developer", "Jun 2024 – Present", "MATMAK", "Dubai, UAE")
    word_note(document, MATMAK_INTRO)
    word_subheading(document, ".NET SaaS platform", before=1.5)
    for text in MATMAK_PLATFORM:
        word_bullet(document, text)
    word_subheading(document, "Delivery, infrastructure, and operations")
    for text in MATMAK_DELIVERY:
        word_bullet(document, text)
    for title, detail, link, bullets in CLIENT_PROJECTS:
        word_subheading(document, "Client project: " + title, detail, link)
        word_bullets(document, bullets)

    for title, dates, company, location, bullets in EARLIER_ROLES:
        word_role(document, title, dates, company, location)
        word_bullets(document, bullets)

    word_section(document, "Education")
    for degree, dates, school, location, detail in EDUCATION:
        word_role(document, degree, dates, school, location)
        word_bullet(document, detail)

    word_section(document, "Certifications and Languages")
    word_skill(document, "Certifications", CERTIFICATIONS)
    word_skill(document, "Languages", LANGUAGES)

    props = document.core_properties
    props.title = "Abdellah Bouhaddi - Full Stack Software Engineer (.NET, React) CV"
    props.subject = "Full stack .NET and React software engineering roles in Dubai and the UAE"
    props.author = "Abdellah Bouhaddi"
    props.keywords = "Full Stack, Software Engineer, .NET, C#, ASP.NET Core, Web API, React, TypeScript, Entity Framework Core, SQL Server, AWS, Docker, CI/CD, Dubai, UAE"
    document.save(DOCX_OUT)


# ---------------------------------------------------------------- PDF

def pdf_styles():
    base = getSampleStyleSheet()["Normal"]

    def style(name, **kw):
        return ParagraphStyle(name, parent=base, **kw)

    return {
        "name": style("Name", fontName="Helvetica-Bold", fontSize=20, leading=22, alignment=TA_CENTER, textColor=PDF_BLACK, spaceAfter=2),
        "title": style("Headline", fontName="Helvetica-Bold", fontSize=10.4, leading=12, alignment=TA_CENTER, textColor=PDF_ACCENT, spaceAfter=4),
        "contact": style("Contact", fontName="Helvetica", fontSize=8.7, leading=10.5, alignment=TA_CENTER, textColor=PDF_BODY),
        "links": style("Links", fontName="Helvetica", fontSize=8.5, leading=10.5, alignment=TA_CENTER, textColor=PDF_ACCENT, spaceAfter=2),
        "h1": style("H1", fontName="Helvetica-Bold", fontSize=10.4, leading=11.5, textColor=PDF_ACCENT, spaceBefore=9, keepWithNext=True),
        "h3": style("H3", fontName="Helvetica-Bold", fontSize=9.7, leading=11.5, textColor=PDF_BLACK, spaceBefore=5, spaceAfter=1, keepWithNext=True),
        "role": style("Role", fontName="Helvetica-Bold", fontSize=10, leading=11.5, textColor=PDF_BLACK),
        "date": style("Date", fontName="Helvetica-Bold", fontSize=8.8, leading=11.5, textColor=PDF_ACCENT, alignment=TA_RIGHT),
        "meta": style("Meta", fontName="Helvetica-Bold", fontSize=9, leading=10.5, textColor=PDF_BODY, spaceAfter=1.2, keepWithNext=True),
        "note": style("Note", fontName="Helvetica-Oblique", fontSize=8.7, leading=10.2, textColor=PDF_MUTED, spaceAfter=0.5, keepWithNext=True),
        "body": style("Body", fontName="Helvetica", fontSize=9.6, leading=12.2, textColor=PDF_BODY, alignment=4),
        "skill": style("Skill", fontName="Helvetica", fontSize=9.3, leading=11.6, textColor=PDF_BODY),
        "bullet": style("Bullet", fontName="Helvetica", fontSize=9.5, leading=11.9, textColor=PDF_BODY, leftIndent=12, bulletIndent=3, spaceAfter=2),
    }


def pdf_footer(canvas, doc):
    canvas.saveState()
    canvas.setFillColor(PDF_MUTED)
    canvas.setFont("Helvetica", 7.2)
    canvas.drawRightString(A4[0] - MARGIN * inch, 0.3 * inch, f"Abdellah Bouhaddi  |  Page {doc.page}")
    canvas.restoreState()


def pdf_section_start(s, text):
    """Section heading kept on the same page as its first block."""
    return [
        Paragraph(escape(text.upper()), s["h1"]),
        HRFlowable(width="100%", thickness=0.7, color=PDF_ACCENT, spaceBefore=1.5, spaceAfter=3.5),
    ]


def pdf_role(s, title, dates, company, location):
    row = Table(
        [[Paragraph(escape(title), s["role"]), Paragraph(escape(dates), s["date"])]],
        colWidths=[PDF_WIDTH * 0.72, PDF_WIDTH * 0.28],
    )
    row.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    meta = Paragraph(f"{escape(company)}&nbsp;&nbsp;|&nbsp;&nbsp;<font name='Helvetica-Oblique' color='#555E66'>{escape(location)}</font>", s["meta"])
    return [row, meta]


def pdf_bullet(s, text):
    return Paragraph(escape(text), s["bullet"], bulletText="•")


def pdf_subheading(s, title, detail=None, link=None, space_before=None):
    html = escape(title)
    if detail:
        html += f"<font name='Helvetica-Oblique' size='8.6' color='#555E66'>&nbsp;&nbsp;|&nbsp;&nbsp;{escape(detail)}</font>"
    if link:
        html += f"<font size='8.6' color='#555E66'>&nbsp;&nbsp;|&nbsp;&nbsp;</font><link href='{link[1]}'><font size='8.6' color='#{ACCENT_HEX}'>{escape(link[0])}</font></link>"
    st = s["h3"] if space_before is None else ParagraphStyle("H3x", parent=s["h3"], spaceBefore=space_before)
    return Paragraph(html, st)


def pdf_skills_table(s, rows):
    data = [[Paragraph(f"<b>{escape(label)}</b>", s["skill"]), Paragraph(escape(text), s["skill"])] for label, text in rows]
    table = Table(data, colWidths=[1.5 * inch, PDF_WIDTH - 1.5 * inch])
    table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 1),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
    ]))
    return table


def build_pdf():
    PDF_OUT.parent.mkdir(parents=True, exist_ok=True)
    s = pdf_styles()
    sep = "&nbsp;&nbsp;|&nbsp;&nbsp;"
    links = sep.join(f"<link href='{url}'>{escape(text)}</link>" for text, url in LINKS)
    story = [
        Paragraph(NAME, s["name"]),
        Paragraph(f"{HEADLINE}<font color='#262D33'>&nbsp;&nbsp;&nbsp;|&nbsp;&nbsp;&nbsp;{escape(HEADLINE_STACK)}</font>", s["title"]),
        Paragraph(f"{LOCATION}{sep}{PHONE}{sep}<link href='mailto:{EMAIL}'><font color='#{ACCENT_HEX}'>{EMAIL}</font></link>", s["contact"]),
        Paragraph(links, s["links"]),
    ]

    story.append(KeepTogether(pdf_section_start(s, "Professional Summary") + [Paragraph(escape(PROFILE), s["body"])]))
    story.append(KeepTogether(pdf_section_start(s, "Technical Skills") + [pdf_skills_table(s, SKILLS)]))

    first = pdf_section_start(s, "Professional Experience")
    first += pdf_role(s, "Full Stack Developer", "Jun 2024 – Present", "MATMAK", "Dubai, UAE")
    first += [Paragraph(escape(MATMAK_INTRO), s["note"]), pdf_subheading(s, ".NET SaaS platform", space_before=1.5), pdf_bullet(s, MATMAK_PLATFORM[0])]
    story.append(KeepTogether(first))
    story += [pdf_bullet(s, t) for t in MATMAK_PLATFORM[1:]]
    story.append(pdf_subheading(s, "Delivery, infrastructure, and operations"))
    story += [pdf_bullet(s, t) for t in MATMAK_DELIVERY]
    for title, detail, link, bullets in CLIENT_PROJECTS:
        story.append(KeepTogether([pdf_subheading(s, "Client project: " + title, detail, link)] + [pdf_bullet(s, t) for t in bullets]))

    for title, dates, company, location, bullets in EARLIER_ROLES:
        story.append(KeepTogether(pdf_role(s, title, dates, company, location) + [pdf_bullet(s, t) for t in bullets]))

    edu = pdf_section_start(s, "Education")
    for i, (degree, dates, school, location, detail) in enumerate(EDUCATION):
        block = pdf_role(s, degree, dates, school, location) + [pdf_bullet(s, detail)]
        story.append(KeepTogether(edu + block if i == 0 else block))

    story.append(KeepTogether(pdf_section_start(s, "Certifications and Languages") + [
        pdf_skills_table(s, [("Certifications", CERTIFICATIONS), ("Languages", LANGUAGES)])
    ]))

    pdf = SimpleDocTemplate(
        str(PDF_OUT),
        pagesize=A4,
        leftMargin=MARGIN * inch,
        rightMargin=MARGIN * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.55 * inch,
        title="Abdellah Bouhaddi - Full Stack Software Engineer (.NET, React) CV",
        author="Abdellah Bouhaddi",
        subject="Full stack .NET and React software engineering roles in Dubai and the UAE",
        keywords="Full Stack, Software Engineer, .NET, C#, ASP.NET Core, React, TypeScript, SQL Server, AWS, Dubai, UAE",
    )
    pdf.build(story, onFirstPage=pdf_footer, onLaterPages=pdf_footer)


if __name__ == "__main__":
    import argparse
    import shutil

    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--version", type=int, help="build this version number instead of incrementing CV_VERSION")
    args = parser.parse_args()

    version = next_version(args.version)
    DOCX_OUT, PDF_OUT = output_paths(version)
    build_docx()
    build_pdf()
    shutil.copyfile(PDF_OUT, SITE_PDF)
    VERSION_FILE.write_text(str(version) + chr(10), encoding="utf-8")
    print(f"CV v{version}")
    print(DOCX_OUT)
    print(PDF_OUT)
    print(SITE_PDF)

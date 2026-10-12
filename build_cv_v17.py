"""Build Abdellah Bouhaddi CV v21 as matching DOCX and PDF deliverables."""

from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import HRFlowable, PageBreak, Paragraph, SimpleDocTemplate, Table, TableStyle


ROOT = Path(__file__).resolve().parent
DOCX_OUT = ROOT / "Abdellah_Bouhaddi_Full_Stack_Software_Engineer_CV_v21.docx"
PDF_OUT = ROOT / "output" / "pdf" / "Abdellah_Bouhaddi_Full_Stack_Software_Engineer_CV_v21.pdf"

FONT = "Aptos"
BLACK = RGBColor(0x00, 0x00, 0x00)
BODY = RGBColor(0x20, 0x27, 0x2B)
MUTED = RGBColor(0x50, 0x59, 0x60)
ACCENT = RGBColor(0x0B, 0x67, 0x63)
PDF_BLACK = colors.HexColor("#000000")
PDF_BODY = colors.HexColor("#20272B")
PDF_MUTED = colors.HexColor("#505960")
PDF_ACCENT = colors.HexColor("#0B6763")

SECTION_COLORS = {
    "Professional Profile": "0B6763",
    "Technical Skills": "0B6763",
    "Professional Experience": "0B6763",
    "Earlier Experience": "0B6763",
    "Education": "0B6763",
    "Certifications and Languages": "0B6763",
}


PROFILE = (
    "Full stack software engineer with 3+ years building and improving business applications in C#, "
    "ASP.NET Core, React, Razor Pages, and SQL. At MATMAK in Dubai, hold end-to-end technical responsibility "
    "across application development, production deployment, servers, databases, infrastructure, integrations, "
    "and ongoing support. Start new client projects from architecture and server provisioning through implementation "
    "and production launch, while also extending the established MATMAK SaaS. Earlier experience includes Java "
    "Spring Boot services, Flutter applications, and computer science instruction."
)

SKILLS = [
    ("Application engineering", "C#, ASP.NET Core, Web API, RESTful APIs, Entity Framework Core, LINQ, dependency injection, JWT and OAuth, Swagger and OpenAPI, OOP, SOLID principles, design patterns, Java, Spring Boot, Hibernate, microservices"),
    ("Frontend and mobile", "React, Razor Pages, JavaScript, TypeScript, HTML5, CSS3, Bootstrap, responsive UI, Flutter, Dart, Android, Firebase"),
    ("Data and delivery", "SQL Server, PostgreSQL, MySQL, MongoDB, Firestore, schema design, query optimization, stored procedures, AWS EC2, Linux, Docker, containerized deployments, Nginx, SSL and TLS, GitHub Actions, Jenkins, CI/CD, Git"),
    ("Integrations and quality", "Stripe, Network International, webhooks, xUnit, JUnit, unit and integration testing, Prometheus, Grafana, Loki, Promtail"),
    ("Additional technologies", "PHP, Laravel, Filament, WordPress, WooCommerce, Angular, Jira, Agile, Scrum, SDLC, code review, secure coding"),
]

MATMAK_CORE = [
    ("End-to-end responsibility", "Operate independently across full-stack development, direct production deployments, production servers and databases, cloud and service accounts, domains and DNS, monitoring, and ongoing support."),
    ("New client delivery", "Start and deliver new client projects from requirements and architecture through server provisioning, full-stack implementation, integrations, production launch, and operation."),
    ("Product engineering", "Maintain and extend an established ASP.NET Core SaaS and PPF cutting platform, building REST APIs, Razor Pages, and workflows for authentication, subscriptions, invoicing, pricing, credits, and cutting operations."),
    ("Solution structure", "Work in a modular .NET solution spanning API, Client, Admin, Models, and Services projects, using Entity Framework Core and SQL-backed application services."),
    ("Frontend and payments", "Redesigned customer and administration interfaces and integrated Stripe one-time payments, recurring subscriptions, invoicing, and webhook-based status updates."),
    ("Performance", "Improved affected EF Core and database read and write operations by about 40% through query tuning and schema redesign."),
    ("Production infrastructure", "Create and configure Linux cloud and VPS servers from purchase through go-live, including operating system setup, SSH access, hardening, firewalls, Nginx reverse proxy, SSL and TLS, application runtimes, and direct production deployment."),
    ("Cloud and reliability", "Manage AWS EC2, production database access, automated backups, and staging; build GitHub Actions workflows and monitor production services with Prometheus, Grafana, Loki, and Promtail."),
    ("Delivery ownership", "Work directly with management and clients on requirements, technical decisions, priorities, production launches, troubleshooting, and ongoing support."),
]

XSHIELD = [
    (
        "Production launch",
        "Delivered the public WooCommerce store at xshield-protect.com end to end, including Linux hosting, domain and DNS, Google Workspace, deployment, and operational support.",
    ),
    (
        "Custom ecommerce development",
        "Built custom PHP pages and plugins, API integrations, and import and export workflows; configured Stripe and Network International payment gateways for UAE checkout.",
    ),
]

HEXIS = [
    ("Platform delivery", "Architected and built a multilingual training application in staging with a Laravel backend, Filament administration, and React public interface."),
    ("Business workflows", "Implemented course planning, locations, registrations, certifications, transactional email, and the Hex'pert quiz flow with emailed access, automatic submission, scoring, and results."),
    ("Reviewed AI assisted workflow", "Used Cursor, OpenAI Codex, custom agent skills, and Model Context Protocol tools where useful, while reviewing and validating changes before release."),
]

BODYFENCE = [
    (
        "Production integration",
        "Developed and enhanced a custom WordPress and PHP plugin for bodyfence.net/hpl, integrating pages, forms, "
        "vehicle patterns, and car kits with a dedicated MATMAK API through custom shortcodes.",
    ),
    (
        "Web administration",
        "Configure domains and subdomains, Google Workspace and professional email, and DNS for production and staging; maintain additional WordPress sites for performance, security, and usability.",
    ),
]

FLUTTER = [
    ("Cross platform development", "Developed Android and iOS applications with Flutter and Dart, responsive Material interfaces, and custom navigation flows."),
    ("Firebase integration", "Integrated Firebase Authentication, Cloud Firestore, Cloud Storage, and push notifications."),
]

SPRING = [
    ("Production services", "Maintained an existing Spring Boot production codebase and REST APIs for the RH Partner HR and payroll SaaS without a greenfield rewrite; deployed each microservice in a separate Docker container."),
    ("Database performance", "Designed schemas and optimized SQL queries, improving affected query performance by 35% and supporting scalability."),
    ("Testing", "Wrote JUnit unit and integration tests with more than 90% coverage in the tested scope, strengthening release quality."),
    ("Production support", "Diagnosed and resolved production defects and supported releases."),
]


def set_word_font(run, size=9.6, bold=False, italic=False, color=BODY):
    run.font.name = FONT
    props = run._element.get_or_add_rPr()
    fonts = props.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        props.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def word_format(paragraph, before=0, after=1.6, line=1.02, keep_next=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.keep_with_next = keep_next
    fmt.widow_control = True


def word_hyperlink(paragraph, text, url, size=8.7):
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{key}"), FONT)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0B6763")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size_el = OxmlElement("w:sz")
    size_el.set(qn("w:val"), str(round(size * 2)))
    props.extend([fonts, color, underline, size_el])
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.extend([props, text_el])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def word_section(document, text):
    p = document.add_paragraph(style="Heading 1")
    word_format(p, before=5.2, after=2.1, line=1.0, keep_next=True)
    hex_color = SECTION_COLORS.get(text, "0B6763")
    border = OxmlElement("w:top")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "8")
    border.set(qn("w:space"), "4")
    border.set(qn("w:color"), hex_color)
    borders = OxmlElement("w:pBdr")
    borders.append(border)
    p._p.get_or_add_pPr().append(borders)
    set_word_font(p.add_run(text.upper()), 10.6, bold=True, color=RGBColor.from_string(hex_color))


def word_subheading(document, text):
    p = document.add_paragraph(style="Heading 2")
    word_format(p, before=3.0, after=1.0, line=1.0, keep_next=True)
    set_word_font(p.add_run(text), 9.6, bold=True, color=BLACK)


def word_role(document, title, dates, company, location):
    p = document.add_paragraph(style="Heading 2")
    word_format(p, before=3.7, after=0.2, line=1.0, keep_next=True)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.15), 2)
    set_word_font(p.add_run(title), 10.1, bold=True, color=BLACK)
    p.add_run("\t")
    set_word_font(p.add_run(dates), 8.8, bold=True, color=ACCENT)
    meta = document.add_paragraph()
    word_format(meta, after=1.0, line=1.0, keep_next=True)
    set_word_font(meta.add_run(company), 9.1, bold=True, color=ACCENT)
    set_word_font(meta.add_run("  |  " + location), 8.9, italic=True, color=MUTED)


def word_bullet(document, label, text):
    p = document.add_paragraph(style="List Bullet")
    word_format(p, after=1.15, line=1.01)
    set_word_font(p.add_run(label + ": "), 9.35, bold=True, color=BLACK)
    set_word_font(p.add_run(text), 9.35)


def word_skill(document, label, text):
    p = document.add_paragraph()
    word_format(p, after=0.8, line=1.0)
    set_word_font(p.add_run(label + ": "), 9.2, bold=True, color=BLACK)
    set_word_font(p.add_run(text), 9.2)


def word_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    word_format(p, after=0, line=1.0)
    set_word_font(p.add_run("ABDELLAH BOUHADDI  |  "), 7.5, color=MUTED)
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText"); instruction.set(qn("xml:space"), "preserve"); instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t"); value.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run = OxmlElement("w:r"); props = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "505960")
    size = OxmlElement("w:sz"); size.set(qn("w:val"), "15")
    props.extend([color, size]); run.extend([props, begin, instruction, separate, value, end]); p._p.append(run)


def configure_word_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(9.6)
    normal.font.color.rgb = BODY
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    for name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = FONT
        style.font.color.rgb = BLACK
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    bullets = document.styles["List Bullet"]
    bullets.font.name = FONT
    bullets.font.size = Pt(9.35)
    bullets._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    bullets._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    bullets.paragraph_format.left_indent = Inches(0.25)
    bullets.paragraph_format.first_line_indent = Inches(-0.14)
    bullets.paragraph_format.tab_stops.add_tab_stop(Inches(0.25))


def build_docx():
    document = Document()
    configure_word_styles(document)
    section = document.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = Inches(0.48); section.bottom_margin = Inches(0.47)
    section.left_margin = Inches(0.62); section.right_margin = Inches(0.62)
    section.header_distance = Inches(0.2); section.footer_distance = Inches(0.22)
    word_page_number(section)

    name = document.add_paragraph(style="Title")
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    word_format(name, after=0.2, line=1.0)
    set_word_font(name.add_run("ABDELLAH BOUHADDI"), 20.5, bold=True, color=BLACK)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    word_format(title, after=1.5, line=1.0)
    set_word_font(title.add_run("FULL STACK SOFTWARE ENGINEER"), 10.7, bold=True, color=BLACK)
    set_word_font(title.add_run("  |  C#  ASP.NET CORE  REACT"), 9.8, bold=True, color=ACCENT)

    contact = document.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    word_format(contact, after=0.5, line=1.0)
    set_word_font(contact.add_run("Dubai, UAE  |  +971 58 534 2499  |  "), 8.9)
    word_hyperlink(contact, "douibouhaddi@gmail.com", "mailto:douibouhaddi@gmail.com", 8.9)

    links_one = document.add_paragraph()
    links_one.alignment = WD_ALIGN_PARAGRAPH.CENTER
    word_format(links_one, after=0.35, line=1.0)
    word_hyperlink(links_one, "abdllahbhd.github.io/Portfolio", "https://abdllahbhd.github.io/Portfolio", 8.6)
    set_word_font(links_one.add_run("  |  "), 8.6, color=MUTED)
    word_hyperlink(links_one, "github.com/AbdllahBhd", "https://github.com/AbdllahBhd", 8.6)

    links_two = document.add_paragraph()
    links_two.alignment = WD_ALIGN_PARAGRAPH.CENTER
    word_format(links_two, after=3.4, line=1.0)
    word_hyperlink(links_two, "linkedin.com/in/abdellah-bouhaddi-1688421a9", "https://www.linkedin.com/in/abdellah-bouhaddi-1688421a9/", 8.6)

    word_section(document, "Professional Profile")
    p = document.add_paragraph()
    word_format(p, after=1.8, line=1.02)
    set_word_font(p.add_run(PROFILE), 9.5)

    word_section(document, "Technical Skills")
    for label, text in SKILLS:
        word_skill(document, label, text)

    word_section(document, "Professional Experience")
    word_role(document, "Full Stack Developer", "June 2024 - Present", "MATMAK", "Dubai, UAE")
    for label, text in MATMAK_CORE:
        word_bullet(document, label, text)
    word_subheading(document, "Selected client delivery at MATMAK")
    word_subheading(document, "XSHIELD  Production")
    for label, text in XSHIELD:
        word_bullet(document, label, text)

    breaker = document.add_paragraph()
    word_format(breaker, after=0, line=1.0)
    breaker.add_run().add_break(WD_BREAK.PAGE)

    word_subheading(document, "HEXIS Training Centre France  Staging")
    for label, text in HEXIS:
        word_bullet(document, label, text)

    word_subheading(document, "HEXIS PRECUT LAB and BODYFENCE  Production")
    for label, text in BODYFENCE:
        word_bullet(document, label, text)

    word_section(document, "Earlier Experience")
    word_role(document, "Mobile Application Developer", "April 2024 - June 2024", "Freelance", "Dubai, UAE")
    for label, text in FLUTTER:
        word_bullet(document, label, text)
    word_role(document, "Backend Developer  Java Spring Boot", "September 2023 - April 2024", "Tech-Instinct", "Bejaia, Algeria")
    for label, text in SPRING:
        word_bullet(document, label, text)
    word_role(document, "Part Time Computer Science Instructor", "September 2022 - June 2023", "Higher School of Computer Science and Technology ESTIN and University Abderrahmane Mira", "Bejaia, Algeria")
    word_bullet(document, "Teaching", "Mentored more than 100 students and created practical labs in programming, algorithms and data structures, computer networks, web development, and mobile development.")

    word_section(document, "Education")
    word_role(document, "Master's Degree in Software Engineering", "2020 - 2022", "University Abderrahmane Mira", "Bejaia, Algeria")
    word_bullet(document, "Thesis", "Built a MERN B2C progressive web ecommerce application with real-time inventory management, payment processing, and a responsive checkout flow.")
    word_role(document, "Bachelor's Degree in Computer Science", "2017 - 2020", "University Abderrahmane Mira", "Bejaia, Algeria")
    word_bullet(document, "Projects", "Built Java inventory management software and desktop C++ and Android Studio game implementations using the Minimax algorithm; participated in the Computer Science Club.")

    word_section(document, "Certifications and Languages")
    word_skill(document, "Certifications", "Java Developer Certificate, January to March 2020; Huawei Certification in Cloud Computing")
    word_skill(document, "Languages", "Kabyle native; Arabic, English, and French fluent")

    document.core_properties.title = "Abdellah Bouhaddi Full Stack Software Engineer CV"
    document.core_properties.subject = "Full stack and .NET software engineering roles in Dubai and the UAE"
    document.core_properties.author = "Abdellah Bouhaddi"
    document.core_properties.keywords = "C#, ASP.NET Core, Web API, REST, React, Entity Framework Core, SQL, AWS, Linux, Docker, CI/CD, Dubai, UAE"
    document.save(DOCX_OUT)


def pdf_styles():
    base = getSampleStyleSheet()
    return {
        "name": ParagraphStyle("Name", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=19.5, leading=20, alignment=TA_CENTER, textColor=PDF_BLACK, spaceAfter=1),
        "title": ParagraphStyle("TitleLine", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=10.1, leading=10.8, alignment=TA_CENTER, textColor=PDF_BLACK, spaceAfter=4),
        "contact": ParagraphStyle("Contact", parent=base["Normal"], fontName="Helvetica", fontSize=8.6, leading=9.4, alignment=TA_CENTER, textColor=PDF_BODY, spaceAfter=2),
        "links": ParagraphStyle("Links", parent=base["Normal"], fontName="Helvetica", fontSize=8.2, leading=9.2, alignment=TA_CENTER, textColor=PDF_ACCENT, spaceAfter=2),
        "links_last": ParagraphStyle("LinksLast", parent=base["Normal"], fontName="Helvetica", fontSize=8.2, leading=9.2, alignment=TA_CENTER, textColor=PDF_ACCENT, spaceAfter=7),
        "h1": ParagraphStyle("H1", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=10.3, leading=10.9, textColor=PDF_BLACK, spaceBefore=0, spaceAfter=2.5, keepWithNext=True),
        "h2": ParagraphStyle("H2", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=9.55, leading=10.15, textColor=PDF_BLACK, spaceBefore=2.5, spaceAfter=1.2, keepWithNext=True),
        "role": ParagraphStyle("Role", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=10.1, leading=10.7, textColor=PDF_BLACK, keepWithNext=True),
        "date": ParagraphStyle("Date", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=8.6, leading=9.5, textColor=PDF_ACCENT, alignment=TA_RIGHT),
        "meta": ParagraphStyle("Meta", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=8.8, leading=9.5, textColor=PDF_ACCENT, spaceAfter=1.5, keepWithNext=True),
        "body": ParagraphStyle("Body", parent=base["Normal"], fontName="Helvetica", fontSize=9.5, leading=10.65, textColor=PDF_BODY, spaceAfter=2),
        "skill": ParagraphStyle("Skill", parent=base["Normal"], fontName="Helvetica", fontSize=9.2, leading=10.0, textColor=PDF_BODY, spaceAfter=1.1),
        "bullet": ParagraphStyle("Bullet", parent=base["Normal"], fontName="Helvetica", fontSize=9.4, leading=10.4, textColor=PDF_BODY, leftIndent=13, firstLineIndent=-8, bulletIndent=4, spaceAfter=1.6),
    }


def pdf_footer(canvas, doc):
    canvas.saveState()
    canvas.setFillColor(PDF_MUTED)
    canvas.setFont("Helvetica", 7.2)
    canvas.drawRightString(letter[0] - 0.62 * inch, 0.27 * inch, f"ABDELLAH BOUHADDI  |  {doc.page}")
    canvas.restoreState()


def pdf_role(story, styles, title, dates, company, location):
    row = Table(
        [[Paragraph(escape(title), styles["role"]), Paragraph(escape(dates), styles["date"])]],
        colWidths=[5.02 * inch, 2.22 * inch],
        hAlign="LEFT",
    )
    row.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    story.append(row)
    story.append(Paragraph(f"{escape(company)}&nbsp;&nbsp; | &nbsp;&nbsp;<font color='#505960'><i>{escape(location)}</i></font>", styles["meta"]))


def pdf_bullets(story, styles, items):
    for label, text in items:
        story.append(Paragraph(f"<b>{escape(label)}:</b> {escape(text)}", styles["bullet"], bulletText="•"))


def pdf_section(story, styles, text):
    hex_color = SECTION_COLORS.get(text, "0B6763")
    section_color = colors.HexColor("#" + hex_color)
    story.append(HRFlowable(width="100%", thickness=0.8, color=section_color, spaceBefore=5, spaceAfter=2))
    story.append(Paragraph(f"<font color='#{hex_color}'>{escape(text.upper())}</font>", styles["h1"]))


def build_pdf():
    PDF_OUT.parent.mkdir(parents=True, exist_ok=True)
    s = pdf_styles()
    story = [
        Paragraph("ABDELLAH BOUHADDI", s["name"]),
        Paragraph("FULL STACK SOFTWARE ENGINEER&nbsp;&nbsp; | &nbsp;&nbsp;<font color='#0B6763'>C#&nbsp;&nbsp; ASP.NET CORE&nbsp;&nbsp; REACT</font>", s["title"]),
        Paragraph("Dubai, UAE&nbsp;&nbsp; | &nbsp;&nbsp;+971 58 534 2499&nbsp;&nbsp; | &nbsp;&nbsp;<link href='mailto:douibouhaddi@gmail.com' color='#0B6763'>douibouhaddi@gmail.com</link>", s["contact"]),
        Paragraph("<link href='https://abdllahbhd.github.io/Portfolio' color='#0B6763'>abdllahbhd.github.io/Portfolio</link>&nbsp;&nbsp; | &nbsp;&nbsp;<link href='https://github.com/AbdllahBhd' color='#0B6763'>github.com/AbdllahBhd</link>", s["links"]),
        Paragraph("<link href='https://www.linkedin.com/in/abdellah-bouhaddi-1688421a9/' color='#0B6763'>linkedin.com/in/abdellah-bouhaddi-1688421a9</link>", s["links_last"]),
    ]
    pdf_section(story, s, "Professional Profile")
    story.append(Paragraph(escape(PROFILE), s["body"]))
    pdf_section(story, s, "Technical Skills")
    for label, text in SKILLS:
        story.append(Paragraph(f"<b>{escape(label)}:</b> {escape(text)}", s["skill"]))
    pdf_section(story, s, "Professional Experience")
    pdf_role(story, s, "Full Stack Developer", "June 2024 - Present", "MATMAK", "Dubai, UAE")
    pdf_bullets(story, s, MATMAK_CORE)
    story.append(Paragraph("Selected client delivery at MATMAK", s["h2"]))
    story.append(Paragraph("XSHIELD  Production", s["h2"]))
    pdf_bullets(story, s, XSHIELD)

    story.append(PageBreak())
    story.append(Paragraph("HEXIS Training Centre France  Staging", s["h2"]))
    pdf_bullets(story, s, HEXIS)

    story.append(Paragraph("HEXIS PRECUT LAB and BODYFENCE  Production", s["h2"]))
    pdf_bullets(story, s, BODYFENCE)

    pdf_section(story, s, "Earlier Experience")
    pdf_role(story, s, "Mobile Application Developer", "April 2024 - June 2024", "Freelance", "Dubai, UAE")
    pdf_bullets(story, s, FLUTTER)
    pdf_role(story, s, "Backend Developer  Java Spring Boot", "September 2023 - April 2024", "Tech-Instinct", "Bejaia, Algeria")
    pdf_bullets(story, s, SPRING)
    pdf_role(story, s, "Part Time Computer Science Instructor", "September 2022 - June 2023", "Higher School of Computer Science and Technology ESTIN and University Abderrahmane Mira", "Bejaia, Algeria")
    pdf_bullets(story, s, [("Teaching", "Mentored more than 100 students and created practical labs in programming, algorithms and data structures, computer networks, web development, and mobile development.")])

    pdf_section(story, s, "Education")
    pdf_role(story, s, "Master's Degree in Software Engineering", "2020 - 2022", "University Abderrahmane Mira", "Bejaia, Algeria")
    pdf_bullets(story, s, [("Thesis", "Built a MERN B2C progressive web ecommerce application with real-time inventory management, payment processing, and a responsive checkout flow.")])
    pdf_role(story, s, "Bachelor's Degree in Computer Science", "2017 - 2020", "University Abderrahmane Mira", "Bejaia, Algeria")
    pdf_bullets(story, s, [("Projects", "Built Java inventory management software and desktop C++ and Android Studio game implementations using the Minimax algorithm; participated in the Computer Science Club.")])

    pdf_section(story, s, "Certifications and Languages")
    story.append(Paragraph("<b>Certifications:</b> Java Developer Certificate, January to March 2020; Huawei Certification in Cloud Computing", s["skill"]))
    story.append(Paragraph("<b>Languages:</b> Kabyle native; Arabic, English, and French fluent", s["skill"]))

    pdf = SimpleDocTemplate(
        str(PDF_OUT),
        pagesize=letter,
        leftMargin=0.62 * inch,
        rightMargin=0.62 * inch,
        topMargin=0.46 * inch,
        bottomMargin=0.45 * inch,
        title="Abdellah Bouhaddi Full Stack Software Engineer CV",
        author="Abdellah Bouhaddi",
        subject="Full stack and .NET software engineering roles in Dubai and the UAE",
    )
    pdf.build(story, onFirstPage=pdf_footer, onLaterPages=pdf_footer)


if __name__ == "__main__":
    build_docx()
    build_pdf()
    print(DOCX_OUT)
    print(PDF_OUT)
